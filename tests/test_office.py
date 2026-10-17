import re
import zipfile
from datetime import date
from pathlib import Path

import docx
import openpyxl
import pytest
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from foldersum.errors import ExtractionError, UnsupportedFileError
from foldersum.extractors import extract_document
from foldersum.extractors.office import OfficeExtractor
from foldersum.models import FileKind
from tests.conftest import SourceFactory

RU_PARAGRAPH = "Договор поставки заключён между сторонами на следующих условиях."


def extract(path: Path, make_source: SourceFactory) -> str:
    return OfficeExtractor().extract(make_source(path, FileKind.OFFICE)).text


# --- docx ---------------------------------------------------------------------------------


def make_docx(path: Path) -> Path:
    document = docx.Document()
    document.add_heading("Quarterly report", level=0)
    document.add_heading("Overview", level=1)
    document.add_paragraph("Sales grew in every region.")
    document.add_paragraph("   ")
    document.add_heading("Details", level=2)
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "North"
    table.cell(1, 1).text = "12"
    document.add_paragraph("Closing remarks.")
    document.save(str(path))
    return path


def test_docx_text_keeps_headings_paragraphs_and_tables_in_order(
    tmp_path: Path, make_source: SourceFactory
) -> None:
    text = extract(make_docx(tmp_path / "report.docx"), make_source)

    assert text == (
        "# Quarterly report\n\n# Overview\n\nSales grew in every region.\n\n"
        "## Details\n\nRegion | Revenue\n\nNorth | 12\n\nClosing remarks."
    )


def test_docx_merged_cells_are_not_repeated(tmp_path: Path, make_source: SourceFactory) -> None:
    document = docx.Document()
    table = document.add_table(rows=1, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "Wide"
    table.cell(0, 2).text = "Narrow"
    path = tmp_path / "merged.docx"
    document.save(str(path))

    assert extract(path, make_source) == "Wide | Narrow"


def test_docx_detects_language(tmp_path: Path, make_source: SourceFactory) -> None:
    document = docx.Document()
    document.add_paragraph(RU_PARAGRAPH)
    path = tmp_path / "ru.docx"
    document.save(str(path))

    assert OfficeExtractor().extract(make_source(path, FileKind.OFFICE)).lang == "ru"


def test_empty_docx_gives_empty_text(tmp_path: Path, make_source: SourceFactory) -> None:
    path = tmp_path / "empty.docx"
    docx.Document().save(str(path))

    assert extract(path, make_source) == ""


# --- xlsx ---------------------------------------------------------------------------------


def make_xlsx(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    sales = workbook.active
    assert sales is not None
    sales.title = "Sales"
    sales.append(["Item", "Qty", "Date"])
    sales.append(["Apples", 3, date(2024, 5, 1)])
    sales.append([None, None, None])
    sales.append(["Pears", 4.5, None])
    notes = workbook.create_sheet("Notes")
    notes.append(["Prepared by finance"])
    workbook.save(path)
    return path


def test_xlsx_sheets_become_sections_of_rows(tmp_path: Path, make_source: SourceFactory) -> None:
    text = extract(make_xlsx(tmp_path / "book.xlsx"), make_source)

    assert text == (
        "## Sales\n\nItem | Qty | Date\n\nApples | 3 | 2024-05-01 00:00:00\n\nPears | 4.5"
        "\n\n## Notes\n\nPrepared by finance"
    )


def test_xlsx_long_sheets_are_truncated(
    tmp_path: Path, make_source: SourceFactory, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr("foldersum.extractors.office._MAX_SHEET_ROWS", 3)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for number in range(10):
        sheet.append([f"row {number}"])
    path = tmp_path / "long.xlsx"
    workbook.save(path)

    text = extract(path, make_source)

    assert "row 2" in text
    assert "row 3" not in text
    assert "[sheet truncated after 3 rows]" in text


# --- pptx ---------------------------------------------------------------------------------


def make_pptx(path: Path) -> Path:
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "Launch plan"
    first.placeholders[1].text = "Ship in May"
    first.notes_slide.notes_text_frame.text = "Mention the budget"
    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "Numbers"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Users"
    table.cell(1, 1).text = "1200"
    group = second.shapes.add_group_shape()
    group.shapes.add_textbox(
        Inches(1), Inches(4), Inches(3), Inches(1)
    ).text_frame.text = "Inside a group"
    third = presentation.slides.add_slide(presentation.slide_layouts[6])
    picture = path.with_suffix(".png")
    Image.new("RGB", (10, 10), "white").save(picture)
    third.shapes.add_picture(str(picture), Inches(1), Inches(1))
    assert third.notes_slide.notes_text_frame is not None  # creates an empty notes page
    presentation.save(str(path))
    return path


def test_pptx_slides_with_text_tables_groups_and_notes(
    tmp_path: Path, make_source: SourceFactory
) -> None:
    text = extract(make_pptx(tmp_path / "deck.pptx"), make_source)

    assert text == (
        "## Slide 1\n\nLaunch plan\n\nShip in May\n\nNotes: Mention the budget\n\n"
        "## Slide 2\n\nNumbers\n\nMetric | Value\n\nUsers | 1200\n\nInside a group\n\n"
        "## Slide 3"
    )


# --- shared behaviour ---------------------------------------------------------------------


def test_dispatcher_routes_office_files(tmp_path: Path, make_source: SourceFactory) -> None:
    path = make_docx(tmp_path / "report.docx")

    doc = extract_document(make_source(path, FileKind.OFFICE))

    assert doc.text.startswith("# Quarterly report")
    assert doc.lang == "en"


def test_long_text_is_truncated_with_a_warning(
    tmp_path: Path, make_source: SourceFactory, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr("foldersum.extractors.office._MAX_CHARS", 40)
    path = make_docx(tmp_path / "report.docx")

    doc = OfficeExtractor().extract(make_source(path, FileKind.OFFICE))

    assert len(doc.text) <= 40
    assert doc.warnings == ("text truncated to the first 40 characters",)


def test_file_that_is_not_a_zip_is_rejected(tmp_path: Path, make_source: SourceFactory) -> None:
    path = tmp_path / "fake.docx"
    path.write_bytes(b"this is not a zip archive")

    with pytest.raises(ExtractionError, match="not a valid Office file"):
        extract(path, make_source)


def test_zip_that_is_not_an_office_document_is_rejected(
    tmp_path: Path, make_source: SourceFactory
) -> None:
    path = tmp_path / "other.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("readme.txt", "hello")

    with pytest.raises(ExtractionError, match=re.escape("Cannot read other.xlsx")):
        extract(path, make_source)


def test_unsupported_suffix_is_rejected(tmp_path: Path, make_source: SourceFactory) -> None:
    path = tmp_path / "legacy.doc"
    path.write_bytes(b"x")

    with pytest.raises(UnsupportedFileError):
        extract(path, make_source)


def test_archive_that_expands_too_much_is_refused(
    tmp_path: Path, make_source: SourceFactory, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr("foldersum.extractors.office._MAX_UNCOMPRESSED_BYTES", 100)
    path = make_docx(tmp_path / "report.docx")

    with pytest.raises(ExtractionError, match="too large"):
        extract(path, make_source)


def test_archive_with_too_many_entries_is_refused(
    tmp_path: Path, make_source: SourceFactory, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr("foldersum.extractors.office._MAX_ARCHIVE_ENTRIES", 2)
    path = make_docx(tmp_path / "report.docx")

    with pytest.raises(ExtractionError, match="too large"):
        extract(path, make_source)


# --- regressions found in review -----------------------------------------------------------


def test_docx_content_controls_are_read(tmp_path: Path, make_source: SourceFactory) -> None:
    document = docx.Document()
    document.add_paragraph("After the cover")
    control = parse_xml(
        f"<w:sdt {nsdecls('w')}><w:sdtContent><w:p><w:r><w:t>Cover title</w:t></w:r></w:p>"
        "</w:sdtContent></w:sdt>"
    )
    document.element.body.insert(0, control)
    path = tmp_path / "sdt.docx"
    document.save(str(path))

    assert extract(path, make_source) == "Cover title\n\nAfter the cover"


def test_docx_nested_tables_follow_their_row(tmp_path: Path, make_source: SourceFactory) -> None:
    document = docx.Document()
    outer = document.add_table(rows=1, cols=2)
    outer.cell(0, 0).text = "Outer"
    inner = outer.cell(0, 1).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "Inner A"
    inner.cell(0, 1).text = "Inner B"
    path = tmp_path / "nested.docx"
    document.save(str(path))

    text = extract(path, make_source)

    assert "Outer" in text
    assert "Inner A | Inner B" in text
    assert text.index("Outer") < text.index("Inner A")


def test_xlsx_with_a_wrong_declared_dimension_is_read_completely(
    tmp_path: Path, make_source: SourceFactory
) -> None:
    source = make_xlsx(tmp_path / "book.xlsx")
    broken = tmp_path / "broken.xlsx"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(broken, "w") as patched:
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                assert b'<dimension ref="A1:C4"/>' in data
                data = data.replace(b'<dimension ref="A1:C4"/>', b'<dimension ref="A1"/>')
            patched.writestr(item, data)

    text = extract(broken, make_source)

    assert "Apples | 3" in text
    assert "Pears | 4.5" in text


def test_real_zip_bomb_is_refused_without_touching_the_limits(
    tmp_path: Path, make_source: SourceFactory
) -> None:
    path = tmp_path / "bomb.docx"
    megabyte = bytes(1024 * 1024)
    with (
        zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive,
        archive.open("word/document.xml", "w", force_zip64=True) as member,
    ):
        for _ in range(120):
            member.write(megabyte)

    assert path.stat().st_size < 1024 * 1024  # tiny on disk, 120 MB when unpacked
    with pytest.raises(ExtractionError, match="too large"):
        extract(path, make_source)
