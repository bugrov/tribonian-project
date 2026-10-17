"""Word, Excel and PowerPoint files (.docx, .xlsx, .pptx) turned into structured plain text."""

import zipfile
from collections.abc import Callable, Iterator
from pathlib import Path

import docx
import openpyxl
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.base import BaseShape

from foldersum.errors import ExtractionError, UnsupportedFileError
from foldersum.langdetect_utils import detect_language
from foldersum.models import ExtractedDoc, SourceFile

# Office files are zip archives that may come from untrusted sources, so refuse zip bombs.
_MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
_MAX_ARCHIVE_ENTRIES = 5000
# Keep the text a model will have to read bounded: a huge sheet is not worth hours of map-reduce.
_MAX_CHARS = 300_000
_MAX_SHEET_ROWS = 2000

_Reader = Callable[[Path], list[str]]


def _check_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
    except (zipfile.BadZipFile, OSError) as exc:
        raise ExtractionError(f"{path.name} is not a valid Office file: {exc}") from exc
    if (
        len(entries) > _MAX_ARCHIVE_ENTRIES
        or sum(entry.file_size for entry in entries) > _MAX_UNCOMPRESSED_BYTES
    ):
        raise ExtractionError(f"{path.name} is too large to read safely")


def _row(cells: list[str]) -> str | None:
    """A table row as one line, or None when every cell is empty."""
    cleaned = [" ".join(cell.split()) for cell in cells]
    while cleaned and not cleaned[-1]:
        cleaned.pop()
    return " | ".join(cleaned) if cleaned else None


def _docx_blocks(
    container: BaseOxmlElement, document: docx.document.Document
) -> Iterator[Paragraph | Table]:
    """Paragraphs and tables in document order, looking inside content controls (cover pages,
    tables of contents and form fields are wrapped in ``w:sdt``)."""
    for child in container.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)
        elif child.tag == qn("w:sdt"):
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                yield from _docx_blocks(content, document)


def _docx_table_lines(table: Table) -> list[str]:
    """Rows as ``a | b`` lines; tables nested in a cell follow the row that contains them."""
    lines = []
    for table_row in table.rows:
        seen: set[int] = set()
        cells = []
        nested: list[str] = []
        for cell in table_row.cells:
            if id(cell._tc) not in seen:  # a merged cell is returned once per column it spans
                seen.add(id(cell._tc))
                cells.append(cell.text)
                for inner in cell.tables:
                    nested.extend(_docx_table_lines(inner))
        line = _row(cells)
        if line:
            lines.append(line)
        lines.extend(nested)
    return lines


def _read_docx(path: Path) -> list[str]:
    """Headings become Markdown headings; tables become ``a | b`` lines."""
    parts: list[str] = []
    document = docx.Document(str(path))
    for block in _docx_blocks(document.element.body, document):
        if isinstance(block, Table):
            parts.extend(_docx_table_lines(block))
            continue
        text = block.text.strip()
        if not text:
            continue
        style = block.style.name if block.style is not None and block.style.name else ""
        level = style.removeprefix("Heading ")
        if style == "Title":
            parts.append(f"# {text}")
        elif style.startswith("Heading ") and level.isdigit():
            parts.append(f"{'#' * min(int(level), 6)} {text}")
        else:
            parts.append(text)
    return parts


def _read_xlsx(path: Path) -> list[str]:
    """Each sheet as a ``## name`` section of ``a | b`` rows (cached formula values)."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            parts.append(f"## {sheet.title}")
            if isinstance(sheet, ReadOnlyWorksheet):
                sheet.reset_dimensions()  # some generators write a wrong or stale <dimension>
            for number, values in enumerate(sheet.iter_rows(values_only=True)):
                if number >= _MAX_SHEET_ROWS:
                    parts.append(f"[sheet truncated after {_MAX_SHEET_ROWS} rows]")
                    break
                line = _row(["" if value is None else str(value) for value in values])
                if line:
                    parts.append(line)
    finally:
        workbook.close()
    return parts


def _pptx_shape_lines(shape: BaseShape) -> list[str]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [line for child in shape.shapes for line in _pptx_shape_lines(child)]  # type: ignore[attr-defined]
    if getattr(shape, "has_table", False):
        rows = shape.table.rows  # type: ignore[attr-defined]
        lines = (_row([cell.text for cell in row.cells]) for row in rows)
        return [line for line in lines if line]
    if getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text.strip()  # type: ignore[attr-defined]
        return [text] if text else []
    return []


def _read_pptx(path: Path) -> list[str]:
    """Each slide as a ``## Slide N`` section with its text, tables and speaker notes."""
    parts: list[str] = []
    for number, slide in enumerate(Presentation(str(path)).slides, 1):
        parts.append(f"## Slide {number}")
        for shape in slide.shapes:
            parts.extend(_pptx_shape_lines(shape))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                parts.append(f"Notes: {notes.text.strip()}")
    return parts


_READERS: dict[str, _Reader] = {".docx": _read_docx, ".xlsx": _read_xlsx, ".pptx": _read_pptx}


class OfficeExtractor:
    def extract(self, source: SourceFile) -> ExtractedDoc:
        reader = _READERS.get(source.path.suffix.lower())
        if reader is None:
            raise UnsupportedFileError(f"Unsupported Office format: {source.path.name}")
        _check_archive(source.path)
        try:
            text = "\n\n".join(reader(source.path))
        except Exception as exc:  # the three libraries raise many unrelated types on bad files
            raise ExtractionError(f"Cannot read {source.path.name}: {exc}") from exc
        warnings: tuple[str, ...] = ()
        if len(text) > _MAX_CHARS:
            text = text[:_MAX_CHARS]
            warnings = (f"text truncated to the first {_MAX_CHARS} characters",)
        return ExtractedDoc(
            source=source, text=text.strip(), lang=detect_language(text), warnings=warnings
        )
